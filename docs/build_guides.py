from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path

OUT = Path(__file__).parent
NAVY = "17324D"
BLUE = "245B83"
TEAL = "087E8B"
PALE = "EAF3F8"
LIGHT = "F4F7FA"
TEXT = "263442"
MUTED = "5F7080"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=TEXT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.66)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 25, "000000"), ("Heading 1", 16, BLUE), ("Heading 2", 11.5, TEAL)]:
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    return doc


def title_block(doc, title, subtitle, strap):
    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(strap.upper())
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, True, "FFFFFF", 8.5)
        shade(t.rows[0].cells[i], NAVY)
    for row_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, TEXT, 8.4)
            if row_idx % 2 == 1:
                shade(cells[i], LIGHT)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def callout(doc, label, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def code(doc, text):
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    shade(cell, "F2F4F7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Inches(0.06)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Cascadia Code"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_source(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def courses_doc():
    d = configure(Document())
    title_block(d, "Courses Backend and MongoDB Guide", "How the Courses page loads recommendations, stores course actions, and connects to MongoDB Atlas", "Developer guide | VIVA learning feature")
    para(d, "This guide explains the complete request path behind Courses and what happens when the configured database connection succeeds or fails. The important operational point is that Atlas must accept the credentials in backend/.env; otherwise this project starts a temporary in-memory database, and course saves do not persist across backend restarts.")
    d.add_heading("How a Courses request flows", level=1)
    table(d, ["Step", "What happens"], [
        ("1. Browser", "Courses.tsx calls the shared request() helper in src/api/api.ts. The helper sends the saved viva_token as a Bearer token."),
        ("2. Express", "The Vite app calls the API at VITE_API_URL, defaulting to http://localhost:4000/api. Express mounts learningRoutes at /api/learning."),
        ("3. Authentication", "requireAuth validates the JWT and sets req.userId. Course recommendations, search, saved items and actions are user-specific."),
        ("4. Controller and services", "coursesController reads the profile and interview history, then calls the course profile and ranking services."),
        ("5. MongoDB", "Mongoose reads Course records and writes interests and CourseActivity records. The response returns course cards to the page."),
    ], [1.2, 5.65])
    d.add_heading("Data stored for the feature", level=1)
    table(d, ["Model", "Purpose"], [
        ("Course", "Catalog item: title, provider, URL, skills, level, price/free flag, quality rank and credential details."),
        ("CourseActivity", "Per-user view, save, complete or dislike action linked to a Course ID."),
        ("User.profile.interests", "Up to 15 interest topics selected on the Courses page."),
        ("Interview", "Completed interview answer evaluations; missingPoints can inform course recommendations."),
        ("ApiCache", "Cached live YouTube search responses when YOUTUBE_API_KEY is configured."),
    ], [1.55, 5.3])
    d.add_heading("Course API routes", level=1)
    table(d, ["Method and route", "Purpose"], [
        ("GET /api/learning/recommendations", "Personalized ranking, saved IDs, interests and recommendation signals."),
        ("GET /api/learning/search?q=...", "Search the catalog; supports free, level and platform filters, plus optional YouTube results."),
        ("GET /api/learning/certificates", "List certificate-bearing courses in credential priority order."),
        ("GET /api/learning/saved", "Return courses saved by the signed-in user."),
        ("POST /api/learning/interests", "Save selected interests to the signed-in user's profile."),
        ("POST /api/learning/courses/:id/:action", "Record view, save, complete or dislike."),
        ("DELETE /api/learning/courses/:id/:action", "Undo save, complete or dislike."),
    ], [3.25, 3.6])
    d.add_heading("How recommendations are ranked", level=1)
    para(d, "The profile service combines selected interests, target role, resume skills and missing points from up to five recent completed interviews. It gives each topic a weight and explanation. MongoDB text search finds matching courses; the recommender blends normalized relevance, catalog quality and level match. Completed and disliked courses are excluded from later recommendations.")
    callout(d, "Certificate ranking", "The Certificates tab puts NPTEL/SWAYAM exam-based academic certificates first, then issuer-backed Coursera professional certificates, then Udemy completion certificates. The ranking is a product label for comparison, not a promise of employer acceptance.")
    d.add_heading("How MongoDB connection and fallback work", level=1)
    para(d, "At startup, backend/server.js loads backend/.env and uses MONGO_URI. If it is absent, the code defaults to mongodb://127.0.0.1:27017/viva_interview. It attempts mongoose.connect with a three-second server-selection timeout.")
    numbered(d, "If the configured URI connects, Mongoose uses that database and the startup log reports its host and database name.")
    numbered(d, "If the connection fails, the server attempts MongoMemoryServer and starts a temporary database named viva_interview. This fallback keeps the API available but stores data only in memory.")
    numbered(d, "After connecting, the server upserts the curated catalog and initializes Course indexes. The curated catalog is reloaded on each backend start.")
    callout(d, "Current project status", "The Atlas URI is configured in backend/.env, but the reported `bad auth` error means Atlas rejected the credentials. The server then used its in-memory fallback. Resolve Atlas authentication before relying on saved interests or course actions.", "FFF3E8")
    d.add_heading("Connect to Atlas", level=1)
    numbered(d, "In Atlas, create or confirm a database user under Security > Database Access. This is a database credential, separate from your Atlas website sign-in.")
    numbered(d, "Check Security > Network Access and allow the computer or network where the backend runs.")
    numbered(d, "Put the Atlas SRV URI in backend/.env and include the database name after the host, for example `...mongodb.net/viva_interview?retryWrites=true&w=majority`. Keep the real username and password private.")
    numbered(d, "If the database password contains URI-reserved characters such as @, #, / or :, percent-encode those characters in the password portion of the URI.")
    numbered(d, "Stop the old backend process, restart it from backend with `npm start`, and inspect the MongoDB startup line. A successful line should name the Atlas host and viva_interview database; an in-memory fallback line means Atlas is still not connected.")
    d.add_heading("Ways to verify the live connection", level=1)
    bullet(d, "Backend terminal: look for `MongoDB connected: <Atlas-host>/viva_interview`. The fallback message explicitly says it connected via In-Memory Server.")
    bullet(d, "Atlas dashboard: after making a Courses request, check the cluster's current connections and activity. A deployed cluster alone does not prove the app connected.")
    bullet(d, "The `/api/health` endpoint reports whether Mongoose is connected, but it does not report the database host. Use the backend startup log to distinguish Atlas from memory fallback.")
    d.add_heading("Course data and import", level=1)
    para(d, "The backend seeds backend/data/curated_courses.json at startup. To add a licensed CSV dataset, place it under backend/data and run the importer from backend, using the same MONGO_URI as the API server:")
    code(d, "node scripts/importCourseDataset.js data\\coursera.csv Coursera")
    para(d, "The importer supports common title and URL header names. It connects directly to MONGO_URI and does not use the server's in-memory fallback, so importing while the API is on fallback would write to a different database.")
    d.add_heading("Relevant project files", level=1)
    table(d, ["File", "Responsibility"], [
        ("backend/server.js", "Loads configuration, connects MongoDB, starts fallback, seeds courses and mounts routes."),
        ("backend/routes/learningRoutes.js", "Registers authenticated learning and course endpoints."),
        ("backend/controllers/coursesController.js", "Validates requests and returns recommendations, searches and activities."),
        ("backend/services/courseProfile.js", "Builds interest, role, skill and interview-gap topics."),
        ("backend/services/courseRecommender.js", "Searches and ranks catalog entries."),
        ("backend/models/Course.js and CourseActivity.js", "Defines catalog and user activity records."),
        ("src/api/api.ts and src/screens/Courses.tsx", "Frontend API calls, filters, tabs and course cards."),
    ], [3.1, 3.75])
    d.save(OUT / "Courses_Backend_and_MongoDB_Guide.docx")


def llm_doc():
    d = configure(Document())
    title_block(d, "LLM Interview Changes Overview", "Implemented model, retrieval, question-generation and evaluation flow for VIVA interviews", "Implementation report | AI interview feature")
    para(d, "The interview AI is split between the Express backend and a Python Flask agent. Express owns authenticated interview sessions and persistence; Flask runs the model workflows. Groq provides chat completions and audio transcription, while sentence-transformers supplies local semantic embeddings for retrieval.")
    d.add_heading("Architecture at a glance", level=1)
    table(d, ["Component", "Role"], [
        ("React client", "Starts interviews, sends answers and audio, and displays question/evaluation responses."),
        ("Express API, default port 4000", "Normalizes setup, plans question count, persists interviews and calls the Python agent."),
        ("Python Flask agent, default port 5100", "Serves question-bank generation, verbal evaluation, coding/SQL evaluation and transcription routes."),
        ("Groq API", "Runs the configured chat model and Whisper transcription model using GROQ_API_KEY."),
        ("Sentence Transformers", "Embeds retrieved web content and interview queries using all-MiniLM-L6-v2 by default."),
        ("MongoDB", "Stores interview setup, question bank, candidate answers, evaluation and completion summary."),
    ], [2.15, 4.7])
    d.add_heading("Interview start and question bank", level=1)
    numbered(d, "The Express controller normalizes the user's role, domain, difficulty, interview types, profile and resume-use setting, then builds a question plan.")
    numbered(d, "It creates an active Interview record and calls Flask `/generate-bank` through interviewLLMService.js.")
    numbered(d, "The Python bank generator requests a structured JSON bank from the configured Groq model. The controller retries brief startup connection failures.")
    numbered(d, "The Node bankGenerator enriches and validates the returned questions, including easy/medium/hard progression and appropriate coding or SQL workspace problems for technical interviews.")
    numbered(d, "If AI bank generation fails, the controller uses the local question generator and stores the resulting bank in the Interview record.")
    d.add_heading("Answer evaluation and adaptive follow-ups", level=1)
    para(d, "The client submits the active question, candidate answer, interview setup, recent history and any coding/SQL submission. Express forwards this payload to Flask `/evaluate` and stores the evaluation in the answer record.")
    table(d, ["Answer section", "Agent behavior"], [
        ("Verbal", "Loads the normal-question agent lazily, retrieves relevant web context for the interview topic and evaluates the answer against the role and question."),
        ("Coding or SQL", "Loads the coding/SQL agent lazily and asks for structured scoring across correctness, reasoning, complexity, quality and edge cases."),
    ], [1.5, 5.35])
    para(d, "The structured evaluation includes a 0–10 score, strengths, missing points, follow-up recommendation and a concise summary. The controller permits at most one focused follow-up per main question, blocks a follow-up after an explicit pass or `I don't know`, and advances or completes the interview according to the plan and model signal.")
    callout(d, "Evaluation boundary", "The coding prompt explicitly says not to claim code was executed or invent test results. The model reviews the submitted answer and any supplied execution information; this agent code does not itself run candidate code.")
    d.add_heading("Verbal retrieval augmented generation", level=1)
    para(d, "For a verbal interview, the agent builds a small knowledge base from pages found through DuckDuckGo, unless the user supplies direct URLs in custom topics. It cleans page text, splits it into chunks, embeds the chunks and stores the vectors in a bounded in-process cache keyed by relevant interview setup.")
    numbered(d, "The question or answer is used to form a retrieval query grounded in the selected domain, subdomain, role and objective.")
    numbered(d, "The agent embeds the query and ranks chunks by cosine similarity. Normalized vectors allow cosine similarity to be computed as a dot product.")
    numbered(d, "The top relevant chunks are added to the Groq prompt as reference context. The agent treats scraped page text as untrusted data and ignores embedded instructions.")
    numbered(d, "If search or page extraction returns no useful context, the prompt instructs the model to use general domain knowledge without claiming that a source supports it.")
    para(d, "The RAG cache lives in Python process memory, is limited to eight setup entries, and resets when the agent restarts. It is not a persistent vector database.")
    d.add_heading("Audio transcription", level=1)
    para(d, "Express accepts the recorded audio and forwards it to Flask `/transcribe`. Flask temporarily writes the uploaded file, sends it to Groq's `whisper-large-v3-turbo` transcription endpoint, returns the transcript, and removes the temporary file afterward.")
    d.add_heading("Configuration and startup", level=1)
    table(d, ["Setting", "Purpose"], [
        ("GROQ_API_KEY", "Required for chat calls and transcription. Keep it in backend/.env; never place it in frontend source."),
        ("INTERVIEW_LLM_MODEL", "Chat model name. Current example default: `openai/gpt-oss-120b`. Set this to a model available to the configured Groq account."),
        ("RAG_EMBED_MODEL", "Local embedding model; default is `all-MiniLM-L6-v2`. First use may download model files."),
        ("RAG_TOP_K / RAG_MAX_PAGES", "Controls retrieved chunks and web pages; defaults are 4 and 4."),
        ("START_INTERVIEW_AGENT", "Set to false to skip Flask startup; otherwise Express spawns backend/python/interview.py."),
        ("INTERVIEW_AGENT_PORT", "Flask port, default 5100. Express calls it at INTERVIEW_AGENT_URL or localhost:5100."),
        ("PYTHON_BIN", "Optional Python executable path. If unset, server tries the project virtual environment, then system Python."),
    ], [2.2, 4.65])
    code(d, "cd backend\n# Add a valid GROQ_API_KEY to backend/.env\nnpm start")
    para(d, "The targeted agent dependencies are listed in backend/python/requirements_interview.txt. The server logs whether Flask becomes ready on port 5100. Flask loads verbal and coding agents only when first used, which keeps initial startup lighter.")
    d.add_heading("Files changed or added for the LLM flow", level=1)
    table(d, ["File", "Change"], [
        ("backend/python/interview.py", "Flask routes for health, transcription, answer evaluation and question-bank generation; lazily loads the agents."),
        ("backend/python/normal_questions_agent.py", "Web retrieval, cleaning, chunking, embeddings, in-memory RAG cache, verbal questions and evaluations."),
        ("backend/python/coding_sql_agent.py", "Separate structured prompts for coding/SQL question generation and evaluation."),
        ("backend/services/interviewLLMService.js", "Node client for Flask endpoints and multipart audio forwarding."),
        ("backend/services/bankGenerator.js", "Local question fallback and technical question-bank enrichment."),
        ("backend/controllers/liveInterviewController.js", "Setup normalization, question planning, retry handling, persistence and adaptive follow-up flow."),
        ("backend/server.js", "Starts the Python interview agent as a child process and waits for its port."),
        ("backend/python/requirements_interview.txt and backend/.env.example", "Document required Python packages and example runtime settings."),
    ], [3.05, 3.8])
    d.add_heading("Operational limitations", level=1)
    bullet(d, "A missing or invalid Groq key prevents model and transcription calls; question-bank generation can fall back locally, but answer evaluation and transcription need the agent/provider to work.")
    bullet(d, "DuckDuckGo results and page access can change or be blocked. RAG may have no retrieved context for a topic.")
    bullet(d, "A slow first request may load/download the embedding model. Ensure the Python environment has the listed dependencies and network access for model retrieval.")
    bullet(d, "The prompts include safeguards, but LLM evaluations can still be imperfect. Review scores and feedback as model-generated assessments.")
    d.add_heading("Code reference map", level=1)
    table(d, ["Layer", "Key files"], [
        ("Express interview controller", "backend/controllers/liveInterviewController.js"),
        ("Model bridge and local bank service", "backend/services/interviewLLMService.js; backend/services/bankGenerator.js"),
        ("Python API and agents", "backend/python/interview.py; normal_questions_agent.py; coding_sql_agent.py"),
        ("Runtime configuration", "backend/server.js; backend/.env.example; backend/python/requirements_interview.txt"),
    ], [2.2, 4.65])
    d.save(OUT / "LLM_Interview_Changes_Overview.docx")


courses_doc()
llm_doc()
print("Created both DOCX guides in", OUT)
